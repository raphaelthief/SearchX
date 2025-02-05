import os, sys
import mimetypes
import PyPDF2
from docx import Document
from PIL import Image, ExifTags
from mutagen import File as MutagenFile
from pymediainfo import MediaInfo
import fitz  # PyMuPDF
from openpyxl import load_workbook  # Handling XLSX metadata
import pefile  # For Windows executable files
import json
import xml.etree.ElementTree as ET
import zipfile
import csv
from pykeepass import PyKeePass
import patoolib
import xlrd
import zipfile
import rarfile
from Dependencies.DS_stored import extract_ds_store_metadata

# colorama
from colorama import init, Fore, Style

init() # Init colorama

def extract_metadata(file_path):
    mime_type, _ = mimetypes.guess_type(file_path)

    if mime_type is None:
        print(f"{Fore.YELLOW}Unable to detect MIME type for {Fore.GREEN}{file_path}{Fore.YELLOW}. Attempting fallback methods...")
        if file_path.endswith(".zip"):
            extract_archive_metadata(file_path)  # ZIP
        elif file_path.endswith(".rar"):
            extract_rar_metadata(file_path)  # RAR
        elif file_path.endswith(".kdbx"):
            extract_kdbx_metadata(file_path)
        elif file_path.endswith(".json"):
            extract_json_metadata(file_path)
        elif file_path.endswith(".xml"):
            extract_xml_metadata(file_path)
        elif file_path.endswith(".md"):
            extract_markdown_metadata(file_path)
        elif file_path.endswith(".csv") or file_path.endswith(".txt"):
            extract_text_metadata(file_path)
        elif file_path.endswith(".xls"):
            extract_xls_metadata(file_path)
        elif file_path.lower().endswith("ds_store"):
            extract_ds_store_metadata(file_path)
        else:
            print(f"{Fore.YELLOW}Unsupported file type for metadata extraction : {Fore.GREEN}{file_path}")
        return

    print(f"{Fore.YELLOW}Detected MIME type : {Fore.GREEN}{mime_type}")

    if mime_type.startswith('application/pdf'):
        extract_pdf_metadata(file_path)
    elif mime_type.startswith('application/vnd.openxmlformats-officedocument.wordprocessingml.document'):
        extract_docx_metadata(file_path)
    elif mime_type.startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
        extract_xlsx_metadata(file_path)
    elif mime_type.startswith('application/vnd.ms-excel'):
        extract_xls_metadata(file_path)  # Handle .xls MIME type
    elif mime_type.startswith('image'):
        extract_image_metadata(file_path)
    elif mime_type.startswith('audio') or mime_type.startswith('video'):
        extract_media_metadata(file_path)
    elif mime_type == 'application/x-msdownload':
        extract_executable_metadata(file_path)
    elif file_path.endswith(".zip"):
        extract_archive_metadata(file_path)  # This is the line we need for ZIP files.
    else:
        print(f"{Fore.YELLOW}File type not supported for MIME metadata extraction : {Fore.GREEN}{mime_type}")


def extract_rar_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- RAR Archive Metadata ---{Fore.GREEN}")
    try:
        with rarfile.RarFile(file_path) as archive:
            print(f"{Fore.YELLOW}Archive Contents :")
            for name in archive.namelist():
                info = archive.getinfo(name)
                print(f"{Fore.YELLOW}File Name         : {Fore.GREEN}{name}")
                print(f"{Fore.YELLOW}  Size            : {Fore.GREEN}{info.file_size} bytes")
                print(f"{Fore.YELLOW}  Compressed Size : {Fore.GREEN}{info.compress_size} bytes")
                print(f"{Fore.YELLOW}  Date Modified   : {Fore.GREEN}{info.date_time}")
                print(f"{Fore.YELLOW}---")
    except Exception as e:
        print(f"{Fore.RED}Error reading RAR archive : {e}")


def extract_archive_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Archive Metadata ---")
    try:
        with zipfile.ZipFile(file_path, 'r') as archive:
            print(f"{Fore.YELLOW}Archive Contents :")
            for name in archive.namelist():
                info = archive.getinfo(name)
                print(f"{Fore.YELLOW}File Name         : {Fore.GREEN}{name}")
                print(f"{Fore.YELLOW}  Size            : {Fore.GREEN}{info.file_size} bytes")
                print(f"{Fore.YELLOW}  Compressed Size : {Fore.GREEN}{info.compress_size} bytes")
                print(f"{Fore.YELLOW}  Date Modified   : {Fore.GREEN}{info.date_time}")
                print(f"{Fore.YELLOW}---")
    except Exception as e:
        print(f"{Fore.RED}Error reading archive file : {e}")


def extract_pdf_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- PDF Detailed Metadata ---")
    try:
        pdf_document = fitz.open(file_path)
        metadata = pdf_document.metadata
        print(f"{Fore.YELLOW}Metadata extracted using PyMuPDF :")
        for key, value in metadata.items():
            print(f"{Fore.YELLOW}{key} : {Fore.GREEN}{value}")

        print(f"{Fore.YELLOW}Number of pages       : {Fore.GREEN}{pdf_document.page_count}")
        print(f"{Fore.YELLOW}Document is encrypted : {Fore.GREEN}{pdf_document.is_encrypted}")

    except Exception as e:
        print(f"{Fore.RED}Error reading PDF metadata : {e}")


def extract_docx_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- DOCX Detailed Metadata ---")
    try:
        doc = Document(file_path)
        core_props = doc.core_properties
        print(f"{Fore.YELLOW}Title              : {Fore.GREEN}", core_props.title)
        print(f"{Fore.YELLOW}Subject            : {Fore.GREEN}", core_props.subject)
        print(f"{Fore.YELLOW}Keywords           : {Fore.GREEN}", core_props.keywords)
        print(f"{Fore.YELLOW}Author             : {Fore.GREEN}", core_props.author)
        print(f"{Fore.YELLOW}Last modified by   : {Fore.GREEN}", core_props.last_modified_by)
        print(f"{Fore.YELLOW}Creation Date      : {Fore.GREEN}", core_props.created)
        print(f"{Fore.YELLOW}Last Modified Date : {Fore.GREEN}", core_props.modified)
        print(f"{Fore.YELLOW}Comments           : {Fore.GREEN}", core_props.comments)
    except Exception as e:
        print(f"{Fore.RED}Error reading the DOCX file : {e}")


def extract_xlsx_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- XLSX Detailed Metadata ---")
    try:
        workbook = load_workbook(file_path)
        props = workbook.properties
        print(f"{Fore.YELLOW}Title         : {Fore.GREEN}{props.title}")
        print(f"{Fore.YELLOW}Author        : {Fore.GREEN}{props.creator}")
        print(f"{Fore.YELLOW}Creation Date : {Fore.GREEN}{props.created}")
        print(f"{Fore.YELLOW}Modified Date : {Fore.GREEN}{props.modified}")
        print(f"{Fore.YELLOW}Company       : {Fore.GREEN}{props.company}")
        print(f"{Fore.YELLOW}Manager       : {Fore.GREEN}{props.manager}")
    except Exception as e:
        print(f"{Fore.RED}Error reading Excel metadata : {e}")


def extract_xls_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- XLS Detailed Metadata ---")
    try:
        workbook = xlrd.open_workbook(file_path)
        print(f"{Fore.YELLOW}Author        : {Fore.GREEN}{workbook.properties.author}")
        print(f"{Fore.YELLOW}Last Modified : {Fore.GREEN}{workbook.properties.last_modified}")
        print(f"{Fore.YELLOW}Created       : {Fore.GREEN}{workbook.properties.created}")
        print(f"{Fore.YELLOW}Sheet Names   : {Fore.GREEN}{workbook.sheet_names()}")
    except Exception as e:
        print(f"{Fore.RED}Error reading XLS metadata : {e}")


def extract_image_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Image Detailed Metadata ---")
    try:
        image = Image.open(file_path)
        exif_data = image._getexif()
        if exif_data:
            for tag_id, value in exif_data.items():
                tag = ExifTags.TAGS.get(tag_id, tag_id)
                print(f"{Fore.YELLOW}{tag} : {Fore.GREEN}{value}")
        else:
            print(f"{Fore.RED}No EXIF metadata found")
    except Exception as e:
        print(f"{Fore.RED}Error reading the image file : {e}")


def extract_media_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Audio/Video Detailed Metadata ---")
    try:
        media_info = MediaInfo.parse(file_path)
        for track in media_info.tracks:
            if track.track_type == "General":
                print(f"{Fore.YELLOW}General Info :")
                print(f"{Fore.YELLOW}Title           : {Fore.GREEN}{track.title}")
                print(f"{Fore.YELLOW}Duration        : {Fore.GREEN}{track.duration} ms")
                print(f"{Fore.YELLOW}File size       : {Fore.GREEN}{track.file_size} bytes")
                print(f"{Fore.YELLOW}Overall bitrate : {Fore.GREEN}{track.overall_bit_rate}")
            if track.track_type == "Video":
                print(f"{Fore.YELLOW}Video Info :")
                print(f"{Fore.YELLOW}Resolution : {Fore.GREEN}{track.width}x{track.height}")
                print(f"{Fore.YELLOW}Frame rate : {Fore.GREEN}{track.frame_rate} fps")
                print(f"{Fore.YELLOW}Codec      : {Fore.GREEN}{track.codec}")
            if track.track_type == "Audio":
                print(f"{Fore.YELLOW}Audio Info :")
                print(f"{Fore.YELLOW}Channels    : {Fore.GREEN}{track.channel_s}")
                print(f"{Fore.YELLOW}Sample rate : {Fore.GREEN}{track.sampling_rate}")
                print(f"{Fore.YELLOW}Bitrate     : {Fore.GREEN}{track.bit_rate}")
    except Exception as e:
        print(f"{Fore.RED}Error reading audio/video metadata : {e}")


def extract_executable_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Windows Executable Metadata ---")
    try:
        pe = pefile.PE(file_path)
        print(f"{Fore.YELLOW}Entry Point Address : {Fore.GREEN}", hex(pe.OPTIONAL_HEADER.AddressOfEntryPoint))
        print(f"{Fore.YELLOW}Image Base          : {Fore.GREEN}", hex(pe.OPTIONAL_HEADER.ImageBase))
        print(f"{Fore.YELLOW}Subsystem           : {Fore.GREEN}", pefile.SUBSYSTEM_TYPE[pe.OPTIONAL_HEADER.Subsystem])
        print(f"{Fore.YELLOW}Machine Type        : {Fore.GREEN}", pefile.MACHINE_TYPE[pe.FILE_HEADER.Machine])
        print(f"{Fore.YELLOW}Number of Sections  : {Fore.GREEN}", pe.FILE_HEADER.NumberOfSections)

        print(f"\n{Fore.YELLOW}Sections Info :")
        for section in pe.sections:
            print(f"{Fore.YELLOW}Name            : {Fore.GREEN}{section.Name.decode().strip()}")
            print(f"{Fore.YELLOW}Virtual Address : {Fore.GREEN}{hex(section.VirtualAddress)}")
            print(f"{Fore.YELLOW}Raw Size        : {Fore.GREEN}{section.SizeOfRawData} bytes")
            print(f"{Fore.YELLOW}Entropy         : {Fore.GREEN}{section.get_entropy()}")
            print(f"{Fore.YELLOW}---")

        pe.close()
    except Exception as e:
        print(f"{Fore.RED}Error reading executable metadata : {e}")


def extract_kdbx_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Keepass KDBX Metadata ---")
    try:
        kp = PyKeePass(file_path)
        print(f"{Fore.YELLOW}Number of entries : {Fore.GREEN}{len(kp.entries)}")
        print(f"{Fore.YELLOW}Number of groups  : {Fore.GREEN}{len(kp.groups)}")
    except Exception as e:
        print(f"{Fore.RED}Error reading Keepass file : {e}")


def extract_archive_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Archive Metadata ---")
    try:
        with zipfile.ZipFile(file_path, 'r') as archive:
            print(f"{Fore.YELLOW}Archive Contents :{Fore.GREEN}")
            for name in archive.namelist():
                print(name)
    except Exception as e:
        print(f"{Fore.RED}Error reading archive file : {e}")


def extract_json_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- JSON Metadata ---")
    try:
        with open(file_path, 'r', encoding='utf-8') as json_file:
            data = json.load(json_file)
            print(f"{Fore.YELLOW}JSON Data :{Fore.GREEN}")
            print(json.dumps(data, indent=4))
    except Exception as e:
        print(f"{Fore.RED}Error reading JSON file : {e}")


def extract_xml_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- XML Metadata ---")
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        print(f"{Fore.YELLOW}Root Tag   : {Fore.GREEN}", root.tag)
        print(f"{Fore.YELLOW}Attributes : {Fore.GREEN}", root.attrib)
        for child in root:
            print(f"{Fore.YELLOW}Tag : {Fore.GREEN}{child.tag}, {Fore.YELLOW}Attributes : {Fore.GREEN}{child.attrib}")
    except Exception as e:
        print(f"{Fore.RED}Error reading XML file : {e}")


def extract_markdown_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Markdown Metadata ---")
    try:
        with open(file_path, 'r', encoding='utf-8') as md_file:
            first_line = md_file.readline()
            if first_line.startswith("---"):
                print(f"{Fore.YELLOW}YAML Metadata Detected{Fore.GREEN}")
                while True:
                    line = md_file.readline()
                    if line.startswith("---"):
                        break
                    print(line.strip())
            else:
                print(f"{Fore.RED}No YAML metadata found in the Markdown file")
    except Exception as e:
        print(f"{Fore.RED}Error reading Markdown file : {e}")


def extract_text_metadata(file_path):
    print(f"\n{Fore.YELLOW}--- Text/CSV Metadata ---")
    try:
        if file_path.endswith(".csv"):
            with open(file_path, 'r', encoding='utf-8') as csv_file:
                reader = csv.reader(csv_file)
                print(f"{Fore.YELLOW}First 5 rows of CSV data :{Fore.GREEN}")
                for i, row in enumerate(reader):
                    if i < 5:
                        print(row)
                    else:
                        break
        else:
            with open(file_path, 'r', encoding='utf-8') as text_file:
                print(f"{Fore.YELLOW}First 5 lines of the file :{Fore.GREEN}")
                for i in range(5):
                    print(text_file.readline().strip())

    except Exception as e:
        print(f"{Fore.RED}Error reading text/CSV file : {e}")


def showme(file_path):
    if os.path.exists(file_path):
        extract_metadata(file_path)
    else:
        print(f"The specified file does not exist")
        
        
